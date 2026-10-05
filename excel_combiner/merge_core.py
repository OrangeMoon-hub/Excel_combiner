"""Headless merge engine shared by GUI and CLI entry points."""

import csv
import os
import tempfile
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional

import openpyxl
from openpyxl.workbook.properties import CalcProperties

from .errors import DecisionRequired, InputError, ProcessingError
from .models import MergeRequest, MergeRunResult


Logger = Optional[Callable[[str], None]]


def _emit(logger: Logger, message: str) -> None:
    if logger is not None:
        logger(message)


def _requires_decision(message):
    raise DecisionRequired(message)


@dataclass
class MergeDecisions:
    """Callbacks supplied by a UI; strict defaults refuse to guess."""

    extra_sheet_map: Callable = lambda filename, sheets, targets: _requires_decision(
        '%s 需要明确的 Sheet 映射: %s' % (filename, ', '.join(sheets)))
    extra_columns: Callable = lambda filename, sheet, header_len, rows: _requires_decision(
        '%s / %s 存在超出表头的数据列' % (filename, sheet))
    no_header: Callable = lambda filename, sheet, column: _requires_decision(
        '%s / %s 存在无列名的数据列' % (filename, sheet))
    discarded_columns: Callable = lambda filename, sheet, columns: _requires_decision(
        '%s / %s 的列不在模板中: %s' % (filename, sheet, ', '.join(columns)))
    duplicate_column: Callable = lambda *args: _requires_decision(
        '存在无法自动判断的同名列，请在 GUI 中确认或调整输入')
    warning: Callable[[str, str], None] = lambda title, message: None


@dataclass
class MergeTaskState:
    exceptions: List[dict] = field(default_factory=list)
    cancelled_tables: List[dict] = field(default_factory=list)


def read_csv(filepath):
    with open(filepath, 'r', encoding='utf-8-sig') as source:
        rows = [row for row in csv.reader(source) if row]
    return {'Sheet1': rows} if rows else {}


def read_xlsx(filepath, logger=None):
    """Read formulas as their last calculated values for position-safe merge."""
    formula_wb = openpyxl.load_workbook(filepath, data_only=False, read_only=True)
    value_wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    try:
        result = {}
        converted = 0
        missing_cache = 0
        for worksheet in formula_wb.worksheets:
            value_ws = value_wb[worksheet.title]
            rows = []
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                values = []
                for column_number, value in enumerate(row, 1):
                    formula_cell = worksheet.cell(row=row_number, column=column_number)
                    if formula_cell.data_type == 'f':
                        converted += 1
                        value = value_ws.cell(row=row_number, column=column_number).value
                        if value is None:
                            missing_cache += 1
                    if value is None:
                        value = ''
                    elif row_number == 1:
                        # Column matching is text based, but data rows must keep
                        # their original Excel types (number, date, boolean, text).
                        value = str(value)
                    values.append(value)
                if any(value != '' for value in values):
                    rows.append(values)
            result[worksheet.title] = rows
        if converted:
            _emit(logger, '  %s: %d 个公式按缓存计算值合并，不复制公式' %
                  (Path(filepath).name, converted))
        if missing_cache:
            _emit(logger, '  警告: %s 有 %d 个公式没有缓存计算值，合并时留空' %
                  (Path(filepath).name, missing_cache))
        return result
    finally:
        formula_wb.close()
        value_wb.close()


def read_table(filepath, logger=None):
    return read_csv(filepath) if str(filepath).lower().endswith('.csv') else read_xlsx(filepath, logger)


def read_big_table(filepath, add_source_column=False, logger=None):
    sheets = read_table(filepath, logger)
    if not sheets:
        raise InputError('模板没有任何 Sheet')
    for sheet_name, data in sheets.items():
        if not data or not data[0]:
            raise InputError('Sheet「%s」表头为空' % sheet_name)
    snapshot = {name: [list(data[0])] for name, data in sheets.items()}
    if add_source_column:
        for sheet_name in snapshot:
            header = snapshot[sheet_name][0]
            if not header or header[0] != '表名':
                snapshot[sheet_name][0] = ['表名'] + header
                _emit(logger, '  Sheet「%s」自动在首列插入“表名”列' % sheet_name)
    _emit(logger, '模板快照已读取: %s' % filepath)
    return snapshot


def _basename_no_ext(filename):
    return os.path.splitext(os.path.basename(filename))[0]


def _column_letter(index):
    return openpyxl.utils.get_column_letter(index + 1)


def _sample_values(values, maximum=5):
    return ','.join(str(value) if value not in ('', None) else '' for value in values[:maximum])


def _resolve_duplicates(small_header, rows, duplicate_names, big_header, big_groups,
                        filename, sheet_name, state, decisions, logger):
    big_to_small = {}
    for column_name, small_indexes in duplicate_names.items():
        big_indexes = big_groups.get(column_name, [])
        if not big_indexes:
            continue
        columns = []
        first_values = None
        all_same = True
        for small_index in small_indexes:
            values = [row[small_index] if small_index < len(row) else '' for row in rows]
            columns.append((small_index, values))
            if first_values is None:
                first_values = values
            elif values != first_values:
                all_same = False
        if all_same:
            for big_index in big_indexes:
                big_to_small[big_index] = small_indexes[0]
            continue
        if len(small_indexes) == len(big_indexes):
            for big_index, small_index in zip(big_indexes, small_indexes):
                big_to_small[big_index] = small_index
            continue
        choices = [(index, values, _column_letter(index)) for index, values in columns]
        for position, big_index in enumerate(big_indexes, 1):
            choice = decisions.duplicate_column(
                column_name, position, len(big_indexes), len(small_indexes), choices,
                filename, sheet_name)
            if choice == 'cancel_table':
                return None, True
            if isinstance(choice, int):
                big_to_small[big_index] = choice
            elif choice == 'ignore':
                big_to_small[big_index] = None
        used = {value for key, value in big_to_small.items() if key in big_indexes}
        for small_index in small_indexes:
            if small_index not in used:
                values = [row[small_index] if small_index < len(row) else '' for row in rows]
                state.exceptions.append({
                    'filename': filename, 'sheet': sheet_name, 'row_num': '-',
                    'col_name': column_name, 'col1_pos': _column_letter(small_index),
                    'col1_val': _sample_values(values) + ' (共%d行)' % len(rows),
                    'col2_pos': '-', 'col2_val': '-', 'exc_type': '同名列未选择',
                    'action': '忽略',
                })
    return big_to_small, False


def process_small_table(filepath, filename, big_snapshot, sheet_mapping=None,
                        add_source_column=False, state=None, decisions=None, logger=None):
    state = state if state is not None else MergeTaskState()
    decisions = decisions if decisions is not None else MergeDecisions()
    try:
        small_sheets = read_table(filepath, logger)
    except Exception as exc:
        raise InputError('读取 %s 失败: %s' % (filename, exc))
    big_names = set(big_snapshot)
    small_names = set(small_sheets)
    mapped_sheets = []
    if sheet_mapping is None:
        mapping = {name: name for name in small_sheets}
        extra_sheets = sorted(small_names - big_names)
        if extra_sheets:
            extra_mapping = decisions.extra_sheet_map(filename, extra_sheets, big_names)
            if extra_mapping is None:
                state.cancelled_tables.append({'filename': filename, 'reason': '用户取消合并'})
                return False, None
            mapping.update(extra_mapping)
    else:
        mapping = dict(sheet_mapping)
    for source_sheet, data in small_sheets.items():
        default_target = source_sheet if source_sheet in big_names else None
        target = mapping[source_sheet] if source_sheet in mapping else default_target
        if target not in big_names:
            _emit(logger, '  %s: Sheet「%s」已跳过' % (filename, source_sheet))
            continue
        mapped_sheets.append((source_sheet, target, data))
    matched = {target for _, target, _ in mapped_sheets}
    if not matched:
        state.cancelled_tables.append({'filename': filename, 'reason': 'Sheet名称不匹配'})
        state.exceptions.append({
            'filename': filename, 'sheet': '-', 'row_num': '-', 'col_name': '-',
            'col1_pos': '-', 'col1_val': '-', 'col2_pos': '-', 'col2_val': '-',
            'exc_type': 'Sheet名称不匹配', 'action': '取消合并',
        })
        decisions.warning('Sheet 不匹配', '小表 %s 的 Sheet 名称与模板不一致，已跳过。' % filename)
        return False, None
    row_cache = {name: [] for name in big_snapshot if name in matched}
    for source_sheet, target_sheet, data in mapped_sheets:
        big_header = big_snapshot[target_sheet][0]
        if not data or not data[0]:
            _emit(logger, '  %s / %s: 无表头，跳过' % (filename, source_sheet))
            continue
        small_header = data[0]
        rows = data[1:]
        if not rows:
            continue
        header_len = len(small_header)
        extra_rows = [(number, len(row), row[header_len:])
                      for number, row in enumerate(rows, 2) if len(row) > header_len]
        if extra_rows:
            if decisions.extra_columns(filename, source_sheet, header_len, extra_rows) == 'cancel_table':
                state.cancelled_tables.append({'filename': filename, 'reason': '用户取消合并'})
                return False, None
            maximum = max(len(extra) for _, _, extra in extra_rows)
            for offset in range(maximum):
                values = [str(extra[offset]) if offset < len(extra) else ''
                          for _, _, extra in extra_rows]
                state.exceptions.append({
                    'filename': filename, 'sheet': source_sheet, 'row_num': '全部行',
                    'col_name': '超表头列数', 'col1_pos': _column_letter(header_len + offset),
                    'col1_val': _sample_values(values) + ' (共%d行)' % len(extra_rows),
                    'col2_pos': '-', 'col2_val': '-', 'exc_type': '超表头列数',
                    'action': '忽略多余列',
                })
        for column_index, header in enumerate(small_header):
            if header != '':
                continue
            if decisions.no_header(filename, source_sheet, column_index) == 'cancel_table':
                state.cancelled_tables.append({'filename': filename, 'reason': '用户取消合并'})
                return False, None
            values = [row[column_index] if column_index < len(row) else '' for row in rows]
            state.exceptions.append({
                'filename': filename, 'sheet': source_sheet, 'row_num': '-',
                'col_name': '数据异常，无列名', 'col1_pos': _column_letter(column_index),
                'col1_val': _sample_values(values) + ' (共%d行)' % len(rows),
                'col2_pos': '-', 'col2_val': '-', 'exc_type': '无列名', 'action': '忽略',
            })
        big_groups = {}
        for index, name in enumerate(big_header):
            big_groups.setdefault(name, []).append(index)
        small_groups = {}
        for index, name in enumerate(small_header):
            if name != '':
                small_groups.setdefault(name, []).append(index)
        discarded = [name for name in small_header if name != '' and name not in big_groups]
        if discarded:
            if decisions.discarded_columns(filename, source_sheet, discarded) == 'cancel_table':
                state.cancelled_tables.append({'filename': filename, 'reason': '用户取消合并'})
                return False, None
            for column_name in discarded:
                small_index = small_header.index(column_name)
                values = [row[small_index] if small_index < len(row) else '' for row in rows]
                state.exceptions.append({
                    'filename': filename, 'sheet': source_sheet, 'row_num': '-',
                    'col_name': column_name, 'col1_pos': _column_letter(small_index),
                    'col1_val': _sample_values(values) + ' (共%d行)' % len(rows),
                    'col2_pos': '-', 'col2_val': '-', 'exc_type': '列名不存在于大表',
                    'action': '忽略并继续',
                })
        duplicates = {name: indexes for name, indexes in small_groups.items()
                      if len(indexes) > 1}
        big_to_small, cancelled = _resolve_duplicates(
            small_header, rows, duplicates, big_header, big_groups,
            filename, source_sheet, state, decisions, logger)
        if cancelled:
            state.cancelled_tables.append({'filename': filename, 'reason': '用户取消合并'})
            return False, None
        for row in rows:
            output_row = []
            for big_index, column_name in enumerate(big_header):
                if add_source_column and column_name == '表名':
                    output_row.append(_basename_no_ext(filename))
                elif big_index in big_to_small:
                    small_index = big_to_small[big_index]
                    output_row.append(row[small_index] if small_index is not None and small_index < len(row) else '')
                elif column_name in small_groups:
                    small_index = small_groups[column_name][0]
                    output_row.append(row[small_index] if small_index < len(row) else '')
                else:
                    output_row.append('')
            row_cache[target_sheet].append(output_row)
    return True, row_cache


def write_result_with_template(template_path, output_path, row_data, snapshot):
    keep_vba = str(template_path).lower().endswith('.xlsm')
    workbook = openpyxl.load_workbook(template_path, data_only=False, keep_vba=keep_vba)
    written = []
    try:
        for sheet_name, rows in (row_data or {}).items():
            if not rows:
                continue
            if sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                header = snapshot.get(sheet_name, [[]])[0]
                for column, value in enumerate(header, 1):
                    worksheet.cell(row=1, column=column).value = value
                column_count = max((len(row) for row in rows), default=0)
                for row_number, values in enumerate(rows, 2):
                    for column, value in enumerate(values, 1):
                        worksheet.cell(row=row_number, column=column).value = value
                old_max = worksheet.max_row
                new_max = 1 + len(rows)
                for row_number in range(new_max + 1, old_max + 1):
                    for column in range(1, column_count + 1):
                        cell = worksheet.cell(row=row_number, column=column)
                        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                            cell.value = None
            else:
                worksheet = workbook.create_sheet(title=sheet_name)
                for values in rows:
                    worksheet.append(values)
            written.append(sheet_name)
        if workbook.calculation is None:
            workbook.calculation = CalcProperties()
        workbook.calculation.calcMode = 'auto'
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output_path)
        return written, list(workbook.sheetnames)
    finally:
        workbook.close()


def _audit_sheets(result, state):
    if state.exceptions:
        header = ['文件名', 'Sheet', '行号', '列名', '列1位置', '列1值',
                  '列2位置(若同名)', '列2值(若同名)', '异常类型', '用户操作']
        result['异常记录'] = [header] + [[
            item.get('filename', ''), item.get('sheet', ''), item.get('row_num', ''),
            item.get('col_name', ''), item.get('col1_pos', ''), item.get('col1_val', ''),
            item.get('col2_pos', ''), item.get('col2_val', ''),
            item.get('exc_type', ''), item.get('action', '')] for item in state.exceptions]
    if state.cancelled_tables:
        result['取消合并的表名'] = [['被取消的表名', '取消原因']] + [
            [item['filename'], item['reason']] for item in state.cancelled_tables]


def run_merge(request: MergeRequest, logger=None) -> MergeRunResult:
    template = Path(request.template)
    inputs = [Path(path) for path in request.inputs]
    output = Path(request.output)
    if not template.is_file():
        raise InputError('模板文件不存在: %s' % template)
    if not inputs:
        raise InputError('至少需要一个 --input 输入文件')
    for path in inputs:
        if not path.is_file():
            raise InputError('输入文件不存在: %s' % path)
    if output.exists() and output.resolve() in [template.resolve()] + [path.resolve() for path in inputs]:
        raise InputError('输出路径不能覆盖模板或输入文件: %s' % output)
    state = MergeTaskState()
    try:
        snapshot = read_big_table(template, request.add_source_column, logger)
        result = OrderedDict((name, []) for name in snapshot)
        for path in inputs:
            mapping = request.sheet_mapping if request.sheet_mapping else None
            ok, rows = process_small_table(
                path, path.name, snapshot, mapping, request.add_source_column,
                state, MergeDecisions(), logger)
            if not ok:
                raise InputError('%s 未产生可合并数据' % path.name)
            for sheet_name, values in rows.items():
                result[sheet_name].extend(values)
        output_sheets = OrderedDict((name, rows) for name, rows in result.items() if rows)
        _audit_sheets(output_sheets, state)
        if not output_sheets:
            raise InputError('没有数据可写入')
        output.parent.mkdir(parents=True, exist_ok=True)
        suffix = output.suffix or '.xlsx'
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=output.stem + '-', suffix=suffix, dir=str(output.parent))
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            written, _ = write_result_with_template(template, temporary, output_sheets, snapshot)
            temporary.replace(output)
        finally:
            if temporary.exists():
                temporary.unlink()
        return MergeRunResult(output, written, state.exceptions, state.cancelled_tables)
    except (InputError, DecisionRequired):
        raise
    except Exception as exc:
        raise ProcessingError('合并失败: %s' % exc)
