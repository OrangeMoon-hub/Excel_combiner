"""File-based data-integrity regressions; only GUI choices are simulated."""
import csv
import importlib.util
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]


def load_script(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


merge = load_script('excel_merge', '合并脚本.py')
split = load_script('excel_split', '拆分脚本.py')


class FileTestCase(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        merge._add_source_column = False
        merge._exceptions.clear()
        merge._cancelled_tables.clear()
        merge._log_lines.clear()
        split._log_lines.clear()
        capture = redirect_stdout(io.StringIO())
        capture.__enter__()
        self.addCleanup(capture.__exit__, None, None, None)
        for name in ('_dup_col_multi_dialog', '_extra_cols_dialog',
                     '_no_header_dialog', '_discarded_cols_dialog',
                     '_extra_sheet_map_dialog'):
            mock = patch.object(merge, name, side_effect=AssertionError('Unexpected dialog: ' + name))
            mock.start()
            self.addCleanup(mock.stop)
        warning = patch.object(merge.messagebox, 'showwarning')
        self.warning = warning.start()
        self.addCleanup(warning.stop)

    def workbook(self, filename, sheets):
        filepath = self.folder / filename
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(filepath)
        wb.close()
        return str(filepath)

    def read_workbook(self, filepath):
        wb = openpyxl.load_workbook(filepath)
        self.addCleanup(wb.close)
        return wb

    def merge_files(self, template, source, mapping):
        snapshot = merge._read_big_table(template)
        ok, rows = merge.process_small_table(source, Path(source).name, snapshot, mapping)
        self.assertTrue(ok)
        output = str(self.folder / 'output.xlsx')
        merge.write_result_with_template(template, output, rows, snapshot)
        return self.read_workbook(output)


class SheetMappingTests(FileTestCase):
    def test_reordered_sheets_to_one_target_are_aligned_before_appending(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名', '金额']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名', '金额'], ['张三', 1000]],
            '二组': [['金额', '姓名'], [2000, '李四']],
        })
        result = self.merge_files(template, source, {'一组': '汇总', '二组': '汇总'})
        self.assertEqual(list(result['汇总'].values),
                         [('姓名', '金额'), ('张三', '1000'), ('李四', '2000')])

    def test_legacy_mapping_dialog_also_preserves_each_header(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名', '金额']]})
        source = self.workbook('source.xlsx', {
            '额外': [['金额', '姓名'], [2000, '李四']],
            '汇总': [['姓名', '金额'], ['张三', 1000]],
        })
        with patch.object(merge, '_extra_sheet_map_dialog', return_value={'额外': '汇总'}):
            result = self.merge_files(template, source, None)
        self.assertEqual(list(result['汇总'].values),
                         [('姓名', '金额'), ('李四', '2000'), ('张三', '1000')])

    def test_different_column_sets_leave_missing_values_blank(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名', '金额', '备注']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名', '金额'], ['张三', 1000]],
            '二组': [['备注', '姓名'], ['补录', '李四']],
        })
        result = self.merge_files(template, source, {'一组': '汇总', '二组': '汇总'})
        self.assertEqual(list(result['汇总'].values)[1:],
                         [('张三', '1000', None), ('李四', None, '补录')])

    def test_extra_column_is_audited_against_original_source_sheet(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名', '金额']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名', '金额'], ['张三', 1000]],
            '二组': [['姓名', '内部备注'], ['李四', '待确认']],
        })
        with patch.object(merge, '_discarded_cols_dialog', return_value='ignore'):
            result = self.merge_files(template, source, {'一组': '汇总', '二组': '汇总'})
        self.assertEqual(list(result['汇总'].values)[2], ('李四', None))
        self.assertEqual([(ex['sheet'], ex['col_name']) for ex in merge._exceptions],
                         [('二组', '内部备注')])

    def test_explicitly_skipped_sheet_is_not_included(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名'], ['张三']], '二组': [['姓名'], ['李四']],
        })
        result = self.merge_files(template, source, {'一组': '汇总', '二组': None})
        self.assertEqual(list(result['汇总'].values), [('姓名',), ('张三',)])

    def test_cancel_later_sheet_discards_entire_source_file(self):
        source = self.workbook('source.xlsx', {
            '一组': [['姓名'], ['张三']], '二组': [['姓名', '额外'], ['李四', 'x']],
        })
        with patch.object(merge, '_discarded_cols_dialog', return_value='cancel_table'):
            ok, rows = merge.process_small_table(source, 'source.xlsx',
                {'汇总': [['姓名']]}, {'一组': '汇总', '二组': '汇总'})
        self.assertFalse(ok)
        self.assertIsNone(rows)

    def test_same_headers_and_multiple_targets_still_work(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名']], '另一表': [['姓名']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名'], ['张三']], '二组': [['姓名'], ['李四']],
            '三组': [['姓名'], ['王五']],
        })
        result = self.merge_files(template, source,
                                  {'一组': '汇总', '二组': '汇总', '三组': '另一表'})
        self.assertEqual(list(result['汇总'].values), [('姓名',), ('张三',), ('李四',)])
        self.assertEqual(list(result['另一表'].values), [('姓名',), ('王五',)])

    def test_source_labels_work_with_multiple_sheets(self):
        merge._add_source_column = True
        template = self.workbook('template.xlsx', {'汇总': [['表名', '姓名', '金额']]})
        source = self.workbook('source.xlsx', {
            '一组': [['姓名', '金额'], ['张三', 1000]],
            '二组': [['金额', '姓名'], [2000, '李四']],
        })
        result = self.merge_files(template, source, {'一组': '汇总', '二组': '汇总'})
        self.assertEqual(list(result['汇总'].values)[1:],
                         [('source', '张三', '1000'), ('source', '李四', '2000')])

    def test_all_sheets_skipped_returns_no_rows(self):
        source = self.workbook('source.xlsx', {'人员': [['姓名'], ['张三']]})
        ok, rows = merge.process_small_table(source, 'source.xlsx',
            {'汇总': [['姓名']]}, {'人员': None})
        self.assertFalse(ok)
        self.assertIsNone(rows)
        self.warning.assert_called_once()

    def test_csv_bom_blank_lines_and_leading_zero_identifier(self):
        template = self.workbook('template.xlsx', {'Sheet1': [['姓名', '工号']]})
        source = self.folder / 'source.csv'
        with source.open('w', encoding='utf-8-sig', newline='') as f:
            csv.writer(f).writerows([['工号', '姓名'], [], ['00123', '张三']])
        result = self.merge_files(template, str(source), {'Sheet1': 'Sheet1'})
        self.assertEqual(list(result['Sheet1'].values), [('姓名', '工号'), ('张三', '00123')])


class TemplateWritingTests(FileTestCase):
    def test_shorter_update_clears_old_tail_without_changing_source_or_styles(self):
        template = self.workbook('template.xlsx', {
            '人员': [['姓名', '状态'], ['张三', '旧'], ['李四', '旧'], ['王五', '旧']],
            '说明': [['保留'], ['=1+1']],
        })
        wb = openpyxl.load_workbook(template)
        wb['人员']['A3'].fill = PatternFill('solid', fgColor='FFFF00')
        wb.save(template)
        wb.close()
        original = Path(template).read_bytes()
        source = self.workbook('source.xlsx', {'人员': [['姓名', '状态'], ['赵六', '新']]})
        result = self.merge_files(template, source, {'人员': '人员'})
        self.assertEqual([result['人员'].cell(r, c).value for r in (3, 4) for c in (1, 2)],
                         [None] * 4)
        self.assertEqual(result['人员']['A2'].value, '赵六')
        self.assertEqual(result['人员']['A3'].fill.fgColor.rgb, '00FFFF00')
        self.assertEqual(result['说明']['A2'].value, '=1+1')
        self.assertEqual(Path(template).read_bytes(), original)

    def test_longer_update_keeps_every_new_row(self):
        template = self.workbook('template.xlsx', {'人员': [['姓名'], ['旧']]})
        source = self.workbook('source.xlsx', {'人员': [['姓名'], ['甲'], ['乙'], ['丙']]})
        result = self.merge_files(template, source, {'人员': '人员'})
        self.assertEqual(list(result['人员'].values), [('姓名',), ('甲',), ('乙',), ('丙',)])

    def test_explicit_none_clears_existing_cell(self):
        template = self.workbook('template.xlsx', {'人员': [['姓名', '备注'], ['张三', '旧']]})
        output = str(self.folder / 'output.xlsx')
        merge.write_result_with_template(template, output, {'人员': [['张三', None]]},
                                         {'人员': [['姓名', '备注']]})
        self.assertIsNone(self.read_workbook(output)['人员']['B2'].value)

    def test_merged_cells_in_old_tail_can_be_cleared_without_losing_merge(self):
        template = self.workbook('template.xlsx', {
            '人员': [['姓名', '备注'], ['张三', '旧'], ['旧合并单元格', None]],
        })
        wb = openpyxl.load_workbook(template)
        wb['人员'].merge_cells('A3:B3')
        wb.save(template)
        wb.close()
        source = self.workbook('source.xlsx', {'人员': [['姓名', '备注'], ['李四', '新']]})
        result = self.merge_files(template, source, {'人员': '人员'})
        self.assertIsNone(result['人员']['A3'].value)
        self.assertIn('A3:B3', result['人员'].merged_cells)


class DuplicateColumnTests(FileTestCase):
    def duplicate_merge(self, small_header, values, big_header, **choice):
        template = self.workbook('template.xlsx', {'人员': [big_header]})
        source = self.workbook('source.xlsx', {'人员': [small_header, values]})
        if choice:
            with patch.object(merge, '_dup_col_multi_dialog', **choice):
                return self.merge_files(template, source, {'人员': '人员'})
        return self.merge_files(template, source, {'人员': '人员'})

    def test_ignore_duplicate_really_leaves_output_blank(self):
        result = self.duplicate_merge(['姓名', '备注', '备注'], ['张三', '同意', '不同意'],
                                      ['姓名', '备注'], return_value='ignore')
        self.assertEqual(result['人员']['A2'].value, '张三')
        self.assertIsNone(result['人员']['B2'].value)
        self.assertEqual({ex['col1_pos'] for ex in merge._exceptions}, {'B', 'C'})

    def test_selected_duplicate_column_is_written(self):
        result = self.duplicate_merge(['姓名', '备注', '备注'], ['张三', '同意', '不同意'],
                                      ['姓名', '备注'], return_value=2)
        self.assertEqual(result['人员']['B2'].value, '不同意')

    def test_first_column_index_zero_is_a_valid_choice(self):
        result = self.duplicate_merge(['备注', '备注', '姓名'], ['同意', '不同意', '张三'],
                                      ['姓名', '备注'], return_value=0)
        self.assertEqual(result['人员']['B2'].value, '同意')

    def test_ignore_and_selection_are_independent_per_target_column(self):
        result = self.duplicate_merge(['姓名', '备注', '备注', '备注'], ['张三', '甲', '乙', '丙'],
                                      ['姓名', '备注', '备注'], side_effect=['ignore', 3])
        self.assertEqual(list(result['人员'].values)[1], ('张三', None, '丙'))

    def test_identical_duplicate_values_fill_all_targets_without_dialog(self):
        result = self.duplicate_merge(['姓名', '备注', '备注'], ['张三', '同意', '同意'],
                                      ['姓名', '备注', '备注'])
        self.assertEqual(list(result['人员'].values)[1], ('张三', '同意', '同意'))

    def test_equal_duplicate_counts_keep_source_order_without_dialog(self):
        result = self.duplicate_merge(['姓名', '备注', '备注'], ['张三', '甲', '乙'],
                                      ['姓名', '备注', '备注'])
        self.assertEqual(list(result['人员'].values)[1], ('张三', '甲', '乙'))

    def test_cancel_duplicate_aborts_file(self):
        source = self.workbook('source.xlsx', {'人员': [['备注', '备注'], ['甲', '乙']]})
        with patch.object(merge, '_dup_col_multi_dialog', return_value='cancel_table'):
            ok, rows = merge.process_small_table(source, 'source.xlsx',
                {'人员': [['备注']]}, {'人员': '人员'})
        self.assertFalse(ok)
        self.assertIsNone(rows)


class SplitGroupingTests(FileTestCase):
    def assert_independent_groups(self, values, rename=False):
        source = self.workbook('source.xlsx', {
            '人员': [['部门', '姓名']] + [[v, '人' + str(i)] for i, v in enumerate(values)],
        })
        original = Path(source).read_bytes()
        results = split.split_tables(source, {'人员': '部门'}, rename_sheet=rename)
        self.assertEqual(len(results), len(values))
        self.assertEqual(len({name.casefold() for name in results}), len(values))
        names = []
        reserved = {'CON', 'PRN', 'AUX', 'NUL'}
        reserved.update('COM%d' % i for i in range(1, 10))
        reserved.update('LPT%d' % i for i in range(1, 10))
        for filename, sheets in results.items():
            self.assertLessEqual(len(filename), 31)
            self.assertNotIn(filename.split('.')[0].upper(), reserved)
            self.assertEqual(len(sheets['人员']), 2)
            names.append(sheets['人员'][1][0])
            # Exercise real output writes, including case-insensitive Windows paths.
            output = self.folder / (filename + '.xlsx')
            if rename:
                sheets = {filename: sheets['人员']}
            split.write_xlsx(str(output), sheets)
            wb = self.read_workbook(output)
            self.assertEqual(wb.worksheets[0]['A2'].value, names[-1])
        self.assertCountEqual(names, ['人' + str(i) for i in range(len(values))])
        self.assertEqual(len(list(self.folder.glob('*.xlsx'))), len(values) + 1)
        self.assertEqual(Path(source).read_bytes(), original)

    def test_replaced_characters_do_not_combine_different_groups(self):
        self.assert_independent_groups(['华东/一区', '华东:一区'])

    def test_truncated_names_remain_separate_and_valid_sheet_names(self):
        self.assert_independent_groups(['长' * 31 + '甲', '长' * 31 + '乙'], rename=True)

    def test_case_only_names_do_not_overwrite_files_on_windows(self):
        self.assert_independent_groups(['Sales', 'sales'])

    def test_existing_numeric_suffix_cannot_collide_with_generated_suffix(self):
        self.assert_independent_groups(['A/B', 'A:B', 'A_B_2', 'A_B'])

    def test_trailing_dots_do_not_combine_groups(self):
        self.assert_independent_groups(['部门', '部门.'])

    def test_windows_reserved_names_are_writable(self):
        self.assert_independent_groups(['CON', '_CON', 'NUL', 'COM1', 'LPT1.txt'])

    def test_same_original_value_across_sheets_goes_to_same_file(self):
        source = self.workbook('source.xlsx', {
            '甲表': [['部门', '姓名'], ['A/B', '张三'], ['A:B', '李四']],
            '乙表': [['部门', '姓名'], ['A:B', '王五'], ['A/B', '赵六']],
        })
        results = split.split_tables(source, {'甲表': '部门', '乙表': '部门'})
        self.assertEqual(len(results), 2)
        people = {tuple((sheets['甲表'][1][0], sheets['乙表'][1][0])) for sheets in results.values()}
        self.assertEqual(people, {('张三', '赵六'), ('李四', '王五')})

    def test_normal_groups_repeat_values_and_empty_values(self):
        source = self.workbook('source.xlsx', {
            '甲表': [['部门', '姓名'], ['研发', '张三'], ['研发', '李四'], ['', '跳过']],
            '乙表': [['部门', '姓名'], ['销售', '王五']],
        })
        results = split.split_tables(source, {'甲表': '部门', '乙表': '部门'})
        self.assertEqual(set(results), {'研发', '销售'})
        self.assertEqual(results['研发']['甲表'], [['姓名'], ['张三'], ['李四']])
        self.assertEqual(results['研发']['乙表'], [['姓名']])
        self.assertEqual(results['销售']['甲表'], [['姓名']])


if __name__ == '__main__':
    unittest.main()
