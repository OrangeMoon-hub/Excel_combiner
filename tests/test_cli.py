"""Black-box contracts for the v1.7 headless core and CLI."""
import ast
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


class CliTestCase(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)

    def workbook(self, filename, sheets):
        path = self.folder / filename
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def run_cli(self, *arguments):
        environment = os.environ.copy()
        environment['PYTHONPATH'] = str(ROOT)
        environment['PYTHONUTF8'] = '1'
        return subprocess.run(
            [sys.executable, '-m', 'excel_combiner', *map(str, arguments)],
            cwd=str(ROOT),
            env=environment,
            text=True,
            encoding='utf-8',
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=30,
        )

    def test_version_is_available_without_starting_gui(self):
        completed = self.run_cli('--version')
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(completed.stdout.strip(), 'Excel Combiner 1.7')

    def test_gui_entry_versions_are_ascii_and_do_not_start_gui(self):
        environment = os.environ.copy()
        environment['PYTHONUTF8'] = '1'
        for script, expected in (
                ('合并脚本.py', 'Excel Combiner Merge 1.7'),
                ('拆分脚本.py', 'Excel Combiner Split 1.7')):
            completed = subprocess.run(
                [sys.executable, str(ROOT / script), '--version'],
                cwd=str(ROOT), env=environment, text=True, encoding='utf-8',
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertEqual(completed.stdout.strip(), expected)

    def test_merge_command_writes_name_aligned_output(self):
        template = self.workbook('template.xlsx', {'人员': [['姓名', '金额']]})
        source = self.workbook('source.xlsx', {'人员': [['金额', '姓名'], [100, '张三']]})
        output = self.folder / 'merged.xlsx'

        completed = self.run_cli(
            'merge', '--template', template, '--input', source,
            '--output', output, '--sheet-map', '人员=人员',
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertTrue(output.exists())
        workbook = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(list(workbook['人员'].values), [('姓名', '金额'), ('张三', '100')])
        self.assertIn(str(output), completed.stdout)

    def test_split_command_writes_one_file_per_group(self):
        source = self.workbook('source.xlsx', {
            '人员': [['部门', '姓名'], ['研发', '张三'], ['销售', '李四']],
        })
        output_dir = self.folder / 'split-output'

        completed = self.run_cli(
            'split', '--source', source, '--output-dir', output_dir,
            '--sheet', '人员=部门',
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(sorted(path.name for path in output_dir.glob('*.xlsx')),
                         ['研发.xlsx', '销售.xlsx'])
        workbook = openpyxl.load_workbook(output_dir / '研发.xlsx', data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(list(workbook['人员'].values), [('姓名',), ('张三',)])

    def test_missing_input_returns_input_error_without_partial_output(self):
        template = self.workbook('template.xlsx', {'人员': [['姓名']]})
        output = self.folder / 'should-not-exist.xlsx'

        completed = self.run_cli(
            'merge', '--template', template,
            '--input', self.folder / 'missing.xlsx', '--output', output,
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('不存在', completed.stderr)
        self.assertFalse(output.exists())

    def test_unmatched_sheet_requires_explicit_mapping(self):
        template = self.workbook('template.xlsx', {'汇总': [['姓名']]})
        source = self.workbook('source.xlsx', {'人员': [['姓名'], ['张三']]})
        output = self.folder / 'should-not-exist.xlsx'

        completed = self.run_cli(
            'merge', '--template', template, '--input', source, '--output', output,
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('Sheet 映射', completed.stderr)
        self.assertFalse(output.exists())

    def test_merge_refuses_unconfirmed_discarded_columns(self):
        template = self.workbook('template.xlsx', {'人员': [['姓名']]})
        source = self.workbook('source.xlsx', {
            '人员': [['姓名', '未配置列'], ['张三', '待确认']],
        })
        output = self.folder / 'should-not-exist.xlsx'

        completed = self.run_cli(
            'merge', '--template', template, '--input', source, '--output', output,
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('列不在模板中', completed.stderr)
        self.assertFalse(output.exists())

    def test_split_missing_input_returns_input_error_without_output_dir(self):
        output_dir = self.folder / 'should-not-exist'

        completed = self.run_cli(
            'split', '--source', self.folder / 'missing.xlsx',
            '--output-dir', output_dir, '--sheet', '人员=部门',
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('不存在', completed.stderr)
        self.assertFalse(output_dir.exists())

    def test_split_refuses_unknown_column_without_partial_files(self):
        source = self.workbook('source.xlsx', {
            '人员': [['部门', '姓名'], ['研发', '张三']],
        })
        output_dir = self.folder / 'should-not-contain-files'

        completed = self.run_cli(
            'split', '--source', source, '--output-dir', output_dir,
            '--sheet', '人员=不存在的列',
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('未找到拆分列', completed.stderr)
        self.assertFalse(output_dir.exists())

    def test_split_can_rename_a_single_output_sheet(self):
        source = self.workbook('source.xlsx', {
            '人员': [['部门', '姓名'], ['研发', '张三']],
        })
        output_dir = self.folder / 'split-output'

        completed = self.run_cli(
            'split', '--source', source, '--output-dir', output_dir,
            '--sheet', '人员=部门', '--rename-sheet',
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        workbook = openpyxl.load_workbook(output_dir / '研发.xlsx', data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ['研发'])

    def test_split_refuses_existing_output_directory(self):
        source = self.workbook('source.xlsx', {
            '人员': [['部门', '姓名'], ['研发', '张三']],
        })
        output_dir = self.folder / 'existing'
        output_dir.mkdir()
        marker = output_dir / 'keep.txt'
        marker.write_text('unchanged', encoding='utf-8')

        completed = self.run_cli(
            'split', '--source', source, '--output-dir', output_dir,
            '--sheet', '人员=部门',
        )

        self.assertEqual(completed.returncode, 3)
        self.assertIn('输出目录已存在', completed.stderr)
        self.assertEqual(marker.read_text(encoding='utf-8'), 'unchanged')
        self.assertEqual(list(output_dir.glob('*.xlsx')), [])

    def test_help_is_available_without_starting_gui(self):
        completed = self.run_cli('--help')
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn('merge', completed.stdout)
        self.assertIn('split', completed.stdout)


class HeadlessCoreStructureTests(unittest.TestCase):
    def test_core_modules_do_not_import_tkinter(self):
        package = ROOT / 'excel_combiner'
        core_files = [package / 'merge_core.py', package / 'split_core.py']
        for path in core_files:
            tree = ast.parse(path.read_text(encoding='utf-8'), filename=str(path))
            imported = {
                alias.name
                for node in ast.walk(tree)
                if isinstance(node, ast.Import)
                for alias in node.names
            }
            imported.update(
                node.module or ''
                for node in ast.walk(tree)
                if isinstance(node, ast.ImportFrom)
            )
            self.assertFalse(any(name == 'tkinter' or name.startswith('tkinter.')
                                 for name in imported), path)

    def test_core_modules_have_no_mutable_task_globals(self):
        package = ROOT / 'excel_combiner'
        for path in (package / 'merge_core.py', package / 'split_core.py'):
            tree = ast.parse(path.read_text(encoding='utf-8'), filename=str(path))
            mutable_assignments = []
            for node in tree.body:
                if isinstance(node, (ast.Assign, ast.AnnAssign)):
                    value = node.value
                    if isinstance(value, (ast.List, ast.Dict, ast.Set,
                                          ast.ListComp, ast.DictComp, ast.SetComp)):
                        mutable_assignments.append(node.lineno)
            self.assertEqual(mutable_assignments, [], path)


if __name__ == '__main__':
    unittest.main()
